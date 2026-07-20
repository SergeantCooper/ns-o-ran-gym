#!/usr/bin/env python3
"""build_ppt.py - generate a presentation (.pptx) covering the whole energy-saving RL
project, embedding the actual result figures. Excludes the reverted masking experiment."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

INK   = RGBColor(0x22, 0x22, 0x22)
ACC   = RGBColor(0x00, 0x72, 0xB2)   # blue accent
GREEN = RGBColor(0x00, 0x79, 0x5c)
MUTE  = RGBColor(0x66, 0x66, 0x66)
SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation(); prs.slide_width = SW; prs.slide_height = SH
BLANK = prs.slide_layouts[6]

def _tb(slide, l, t, w, h):
    tf = slide.shapes.add_textbox(l, t, w, h).text_frame
    tf.word_wrap = True
    return tf

def title_bar(slide, title, sub=None):
    tf = _tb(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0))
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = INK
    # accent underline
    ln = slide.shapes.add_shape(1, Inches(0.62), Inches(1.28), Inches(3.2), Pt(4))
    ln.fill.solid(); ln.fill.fore_color.rgb = ACC; ln.line.fill.background()
    if sub:
        tf2 = _tb(slide, Inches(0.62), Inches(1.4), Inches(12.0), Inches(0.5))
        r2 = tf2.paragraphs[0].add_run(); r2.text = sub
        r2.font.size = Pt(14); r2.font.italic = True; r2.font.color.rgb = MUTE

def bullets(slide, items, top=1.9, width=12.1, size=18, left=0.7):
    tf = _tb(slide, Inches(left), Inches(top), Inches(width), Inches(5.2))
    for i, (txt, lvl) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        r = p.add_run(); r.text = ("• " if lvl == 0 else "– ") + txt
        r.font.size = Pt(size - lvl * 3); r.font.color.rgb = INK if lvl == 0 else MUTE
        p.space_after = Pt(7)

def image(slide, path, top=1.7, max_w=11.0, caption=None):
    if not os.path.exists(path):
        return
    w = Inches(max_w); left = Inches((13.333 - max_w) / 2)
    pic = slide.shapes.add_picture(path, left, Inches(top), width=w)
    if caption:
        tf = _tb(slide, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.4))
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = caption; r.font.size = Pt(11); r.font.italic = True; r.font.color.rgb = MUTE

def new(): return prs.slides.add_slide(BLANK)

# ---- 1. Title ----
s = new()
tf = _tb(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(2.0))
p = tf.paragraphs[0]; r = p.add_run(); r.text = "Energy Saving in O-RAN with Reinforcement Learning"
r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = ACC
p2 = tf.add_paragraph(); r2 = p2.add_run()
r2.text = "A digital-twin study: learning to sleep 5G cells without hurting service quality"
r2.font.size = Pt(18); r2.font.color.rgb = INK
p3 = tf.add_paragraph(); r3 = p3.add_run()
r3.text = "ns-3 mmWave O-RAN  ·  imitation + PPO  ·  energy-vs-QoS tradeoff"
r3.font.size = Pt(14); r3.font.italic = True; r3.font.color.rgb = MUTE

# ---- 2. Problem & goal ----
s = new(); title_bar(s, "The problem & the goal")
bullets(s, [
 ("5G base stations (gNBs) draw ~600 W even when almost idle — a large, avoidable energy cost.", 0),
 ("Idea: an intelligent controller that puts unused gNBs to sleep, waking them when demand returns.", 0),
 ("The tension: sleep too aggressively and you drop calls / lose throughput (QoS).", 0),
 ("Company goal: a reinforcement-learning (RL) controller that beats the hand-coded heuristic on the", 0),
 ("energy-vs-QoS tradeoff — more energy saved without hurting throughput or reliability.", 1),
])

# ---- 3. The digital twin ----
s = new(); title_bar(s, "The digital twin (what we simulate)")
bullets(s, [
 ("ns-3 mmWave O-RAN simulator, scenario-three: 1 LTE anchor cell + 7 mmWave gNBs (can sleep), 6 users.", 0),
 ("Traffic redesigned to 3 sharp bursts per episode (+ an always-on baseline user) — so the controller", 0),
 ("must adapt in real time (sleep in the lulls, wake for the bursts).", 1),
 ("Energy model: each ON gNB = 600 W static + 400 W x utilisation; a sleeping gNB ~ 0 W.", 0),
 ("Key constraint: the simulator runs ~15-20 s per 100 ms control step (CPU-bound) — this shapes everything.", 0),
])

# ---- 4. How it works ----
s = new(); title_bar(s, "How it works — real-time closed-loop control")
bullets(s, [
 ("Every 100 ms, a handshake over shared memory (files + semaphores):", 0),
 ("ns-3 emits per-cell KPMs (load, throughput, dropped calls) -> a 61-number observation.", 1),
 ("the controller picks a 7-bit ON/OFF action (one bit per gNB; anchor always ON).", 1),
 ("the action is written back; ns-3 applies it and advances 100 ms.", 1),
 ("Reward the RL optimises:  throughput(Mbps) - energy(kW) - 2 x dropped-calls(RLF).", 0),
 ("The controller plays the role of an O-RAN 'xApp'; we drive the sim directly (no live RIC needed).", 0),
])

# ---- 5. Controllers ----
s = new(); title_bar(s, "The controllers we compared")
bullets(s, [
 ("Heuristic (TwinHeuristic) - the company baseline: load-adaptive, sleeps idle cells, wakes on neighbour load.", 0),
 ("Behaviour Cloning (BC) - a neural net trained to copy the heuristic (reached 100% match).", 0),
 ("PPO (reinforcement learning) - warm-started from BC, then fine-tuned on the reward. The one that can surpass.", 0),
 ("Pruning controller - an aggressive energy-saver with a tunable 'grace' dial (energy-vs-reliability).", 0),
 ("(ns-3's own built-in heuristic is quota-driven and can't adapt to load - hence the custom one.)", 0),
])

# ---- 6. Pipeline ----
s = new(); title_bar(s, "The learning pipeline")
bullets(s, [
 ("1)  Collect expert demonstrations - run the heuristic on the twin, log (observation, action) pairs.", 0),
 ("2)  Behaviour Cloning (+ critic warm-up) - clone the heuristic into the policy AND pre-train its value", 0),
 ("estimator; the warm-up fixes a classic BC->PPO collapse we hit and diagnosed.", 1),
 ("3)  PPO fine-tune (~4 h, live ns-3 rollouts) - improve on the true reward.", 0),
 ("4)  Evaluate deterministically vs the heuristic across 4 seeds; score energy / throughput / RLF.", 0),
 ("All reproducible; trained model + demos committed to the repo.", 0),
])

# ---- 7. Result 1: tradeoff & the tie ----
s = new(); title_bar(s, "Result 1 — the energy-vs-QoS tradeoff", "Heuristic sits near the efficient frontier; plain-reward PPO ties it byte-for-byte")
image(s, "summary_tradeoff.png", top=1.75, max_w=10.5,
      caption="4-seed averages. PPO reproduces the heuristic (a tie). Aggressive pruning saves more energy but costs reliability.")

# ---- 8. Result 2: reward shaping ----
s = new(); title_bar(s, "Result 2 — an improved reward moved RL off the tie",
                      "Potential-based reward shaping -> a distinct, greener + more-reliable policy")
image(s, "figures/2_rl_vs_heuristic_by_metric.png", top=1.8, max_w=11.5,
      caption="Shaped-reward RL vs balanced heuristic (4-seed avg): +6 pts energy AND ~26% fewer dropped calls, ~9% less throughput.")

# ---- 9. Why not a strict win ----
s = new(); title_bar(s, "Why RL doesn't strictly dominate (the key insight)")
bullets(s, [
 ("The heuristic is genuinely near-optimal here - it's hard to beat because it's already good.", 0),
 ("RL is sample-starved: the slow sim affords only ~640 trial-steps, and random on/off exploration", 0),
 ("sleeps idle AND busy cells together -> can't isolate 'sleeping THIS idle cell was the good move'.", 1),
 ("The 43% 'idle-but-powered' waste is real in hindsight, but NOT free to reclaim: sleeping a", 0),
 ("momentarily-idle cell that's needed next causes a dropped call. The heuristic's caution is justified.", 1),
 ("So it's a favourable TRADEOFF (greener + more reliable, slightly less throughput), not a clean sweep.", 0),
])

# ---- 10. Pruning frontier ----
s = new(); title_bar(s, "Operational value — a tunable energy/QoS frontier",
                      "One dial lets an operator choose the operating point; the heuristic is a single fixed point")
image(s, "figures/3_pruning_frontier.png", top=1.8, max_w=9.5,
      caption="Pruning 'grace' dial: from the heuristic's point up to ~57% energy saved, trading throughput as you go.")

# ---- 11. Behaviour over time ----
s = new(); title_bar(s, "What the controller does over time",
                      "Cells sleep through the lulls and wake for the bursts")
image(s, "figures/4_controller_behavior_over_time.png", top=1.8, max_w=10.5,
      caption="One shaped-RL run (seed 999, ~43% saved). RL's 4-seed AVERAGE is ~40% energy saved; "
              "per-seed it varies 25-46%. Idle cells sleep; busy cells stay on.")

# ---- 12. Bottom line ----
s = new(); title_bar(s, "Bottom line & next steps")
bullets(s, [
 ("Delivered: a realistic bursty twin, a working imitation+RL pipeline (BC 100% match, stable PPO),", 0),
 ("a rigorous energy-vs-QoS tradeoff, and a tunable energy-saving controller.", 1),
 ("RL result: with reward shaping, RL learns a greener + more-reliable operating point than the heuristic", 0),
 ("- better on energy AND dropped-calls, slightly lower throughput: a favourable, honest tradeoff.", 1),
 ("It does not STRICTLY beat a strong heuristic on all three axes - the real blocker is the slow", 0),
 ("simulator's tiny sample budget (~640 RL trial-steps), too little to out-tune a good heuristic.", 1),
 ("Next steps: (1) sample-efficient RL - offline RL on logged data, or a fast learned surrogate of the", 0),
 ("sim - to escape that budget;  (2) richer / larger network scenarios;  (3) ship the tunable controller.", 1),
])

# ---- 13. Reproducibility ----
s = new(); title_bar(s, "Reproducibility & deliverables")
bullets(s, [
 ("Two git repos: ns-3-mmwave-oran (the twin) + ns-o-ran-gym (gym / RL / analysis).", 0),
 ("Committed: trained model + demos + obs-stats; all scripts; figures computed live from run folders.", 0),
 ("Docs (self-contained): OVERVIEW.md (how it all works), RESULTS.md (findings & numbers),", 0),
 ("REPRODUCE.md (build & run from scratch), CONTEXT.md (full-journey primer).", 1),
 ("Figures: energy-vs-QoS tradeoff, per-metric comparison, pruning frontier, controller-over-time.", 0),
])

out = "energy_saving_RL_project.pptx"
prs.save(out)
print(f"saved: {out}  ({len(prs.slides._sldIdLst)} slides)")
